"""Land exports of different years side by side.

Layer: tests (integration)

The loader assumes no year: it lands whatever an export holds and lets the rows
say which period that is. These tests build small QAD-like exports with
openpyxl - the current period, an older one, and one delivered later - put them
side by side as the plant would, and land them in a real database.

They need no file from the business, so unlike the acceptance tests in
``test_equs_bronze_load.py`` they also run where the real workbooks are absent.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date
from typing import TYPE_CHECKING

import pytest
from openpyxl import Workbook
from sqlalchemy import text

from energy.adapters.calamine_workbook import CalamineWorkbookReader
from energy.adapters.local_storage import LocalObjectStorage
from energy.adapters.postgres_bronze_repo import PostgresBronzeLandingRepository
from energy.pipelines.flows.equs_bronze_load import LoadAllResult, load_all
from energy.services.bronze_loader import BronzeLoader

if TYPE_CHECKING:
    from pathlib import Path

    from sqlalchemy.engine import Engine

pytestmark = pytest.mark.integration

BUCKET = "bronze-raw"

CURRENT = "Calculo de equivalente - Explorar transacciones por art.xlsx"
OLDER = "Calculo de equivalente - Explorar transacciones por art FY24.xlsx"
LATER = "Calculo de equivalente - Explorar transacciones por art FY23.xlsx"
# An invoices export that only exists for an older period.
INVOICES = "Calculo de equivalente - Unidades facturadas 2024.xlsx"

RECEIPTS = {
    CURRENT: [
        [17240425, date(2025, 12, 10), 70000],
        [17240418, date(2025, 10, 1), 35000],
        [17990001, date(2026, 9, 30), 120],
    ],
    OLDER: [
        [9000001, date(2023, 10, 2), 10],
        [9000002, date(2024, 9, 30), 12],
    ],
    LATER: [
        [8000001, date(2022, 11, 15), 5],
        [8000002, date(2023, 3, 1), 6],
        [8000003, date(2023, 9, 29), 7],
        [8000004, date(2022, 10, 3), 8],
    ],
}
PERIODS = {
    CURRENT: (date(2025, 10, 1), date(2026, 9, 30)),
    OLDER: (date(2023, 10, 2), date(2024, 9, 30)),
    LATER: (date(2022, 10, 3), date(2023, 9, 29)),
    INVOICES: (date(2024, 1, 5), date(2024, 6, 30)),
}


def _export(key_header: str, rows: list[list[object]]) -> bytes:
    """Build an export the way QAD delivers one: data, conditions and info.

    Every export states the same search conditions, October to December 2025,
    which is true of none of them. That is the defect of the real transactions
    export, reproduced on purpose.
    """
    workbook = Workbook()

    datos = workbook.active
    datos.title = "Datos"
    datos.append([key_header, "Fecha Efectiva", "Cantidad"])
    for row in rows:
        datos.append(row)

    conditions = workbook.create_sheet("Condiciones de búsqueda")
    conditions.append(["Campo", "Operador", "Valor", "Monto Hasta"])
    conditions.append(["Fecha Efectiva", "rango", date(2025, 10, 1), date(2025, 12, 31)])

    info = workbook.create_sheet("Info")
    info.append(["Nombre", "Valor"])
    info.append(["Programa Exploración", "icbr011.p"])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@dataclass(frozen=True)
class Landed:
    """One shared load of a folder holding exports of several periods."""

    engine: Engine
    loader: BronzeLoader
    directory: Path
    summary: LoadAllResult


@pytest.fixture(scope="class")
def landed(bronze_engine: Engine, tmp_path_factory: pytest.TempPathFactory) -> Landed:
    """Put three exports side by side and land the folder once."""
    directory = tmp_path_factory.mktemp("raw_excels")
    for name in (CURRENT, OLDER):
        (directory / name).write_bytes(_export("Número de transacción", RECEIPTS[name]))
    (directory / INVOICES).write_bytes(
        _export("Núm Fact", [["F-1", date(2024, 6, 30), 2], ["F-2", date(2024, 1, 5), 4]])
    )

    loader = BronzeLoader(
        reader=CalamineWorkbookReader(),
        storage=LocalObjectStorage(tmp_path_factory.mktemp("bronze_archive")),
        repository=PostgresBronzeLandingRepository(bronze_engine),
        bucket=BUCKET,
    )
    return Landed(bronze_engine, loader, directory, load_all(loader, directory))


def _runs(engine: Engine) -> dict[str, tuple[date | None, date | None, int]]:
    """Return period and row count of every completed run, by file name."""
    with engine.connect() as connection:
        rows = connection.execute(
            text(
                "SELECT parameters ->> 'filename' AS filename, window_start, window_end, "
                "row_count FROM bronze.ingestion_runs WHERE state = 'completed'"
            )
        ).all()
    return {row.filename: (row.window_start, row.window_end, row.row_count) for row in rows}


class TestExportsOfSeveralYears:
    """Three exports in one folder, landed by one load-all with no year given."""

    def test_every_export_is_landed_and_nothing_else_stops_it(self, landed: Landed) -> None:
        """The exports land although the other workbooks of the close are absent."""
        statuses = {
            outcome.path.name: outcome.status
            for outcome in landed.summary.outcomes
            if outcome.path is not None
        }

        assert statuses == {CURRENT: "landed", OLDER: "landed", INVOICES: "landed"}
        assert landed.summary.rows_landed == 7
        assert landed.summary.unmatched == ()
        # The structures and catalogues are not in this folder, and that is said.
        assert not landed.summary.ok

    def test_each_run_covers_the_dates_its_own_rows_hold(self, landed: Landed) -> None:
        """Nobody stated a year: each run reports the first and last date it landed."""
        runs = _runs(landed.engine)

        assert runs == {
            CURRENT: (*PERIODS[CURRENT], 3),
            OLDER: (*PERIODS[OLDER], 2),
            INVOICES: (*PERIODS[INVOICES], 2),
        }

    def test_rows_of_every_year_are_in_bronze(self, landed: Landed) -> None:
        """Receipts of both exports share one table, each pointing at its own file."""
        with landed.engine.connect() as connection:
            rows = connection.execute(
                text(
                    "SELECT f.filename, left(r.payload ->> 'Fecha Efectiva', 4) AS year "
                    "FROM bronze.raw_rct_wo r JOIN bronze.raw_files f USING (raw_file_id)"
                )
            ).all()

        assert len(rows) == 5
        assert {(row.filename, row.year) for row in rows} == {
            (CURRENT, "2025"),
            (CURRENT, "2026"),
            (OLDER, "2023"),
            (OLDER, "2024"),
        }

    def test_what_an_export_states_is_stored_but_not_believed(self, landed: Landed) -> None:
        """The stale conditions are kept as delivered; the period comes from the rows."""
        with landed.engine.connect() as connection:
            stated = connection.execute(
                text(
                    "SELECT search_conditions -> 0 ->> 'monto hasta' "
                    "FROM bronze.ingestion_runs WHERE parameters ->> 'filename' = :filename"
                ),
                {"filename": OLDER},
            ).scalar_one()

        assert stated == "2025-12-31"
        assert _runs(landed.engine)[OLDER][:2] == PERIODS[OLDER]

    def test_a_second_load_all_lands_nothing(self, landed: Landed) -> None:
        """Every export is recognised by its bytes and skipped."""
        before = _runs(landed.engine)

        again = load_all(landed.loader, landed.directory)

        assert again.rows_landed == 0
        assert {o.status for o in again.outcomes if o.path is not None} == {"skipped"}
        assert _runs(landed.engine) == before

    def test_an_export_delivered_later_is_the_only_one_landed(self, landed: Landed) -> None:
        """A new file beside the others brings its rows and leaves theirs alone."""
        (landed.directory / LATER).write_bytes(_export("Número de transacción", RECEIPTS[LATER]))

        again = load_all(landed.loader, landed.directory)

        statuses = {o.path.name: o.status for o in again.outcomes if o.path is not None}
        assert statuses == {
            CURRENT: "skipped",
            LATER: "landed",
            OLDER: "skipped",
            INVOICES: "skipped",
        }
        assert again.rows_landed == 4
        assert _runs(landed.engine)[LATER] == (*PERIODS[LATER], 4)
        with landed.engine.connect() as connection:
            receipts = connection.execute(text("SELECT count(*) FROM bronze.raw_rct_wo")).scalar()
        assert receipts == 9
