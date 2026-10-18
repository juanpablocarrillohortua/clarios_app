"""Tests for the energy-EQ cross service and the SCADA export it reads.

Layer: tests (unit)

The service is given fakes of its two ports: an energy source that answers a
fixed area energy, and an EQ reader that answers a fixed series and records how
it was asked. The SCADA adapter reads a small export built with openpyxl, laid
out as the real one is: a ``t_stamp`` column of naive local times and one Wh
counter per rectifier.
"""

from __future__ import annotations

import io
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import TYPE_CHECKING, Any
from zoneinfo import ZoneInfo

import pytest
from openpyxl import Workbook

from energy.adapters.scada_workbook import CounterWorkbook, ScadaCounterSource
from energy.domain.energy_cross import AreaEnergy, EnergySpan
from energy.domain.equs_gold import DateRange, VersionMonth
from energy.domain.equs_series import BatteryEquivalents, ReconciledAmount, SeriesRow
from energy.domain.exceptions import (
    ConfigurationError,
    InvalidPeriodError,
    SourceUnavailableError,
    UnsupportedFileFormatError,
)
from energy.ports.energy_cross import AreaEnergySource, EquivalentsReader
from energy.services.energy_cross import EnergyCross
from energy.services.equs_gold import BatterySeries

if TYPE_CHECKING:
    from collections.abc import Sequence
    from pathlib import Path

    from energy.domain.equs_series import Frequency

pytestmark = pytest.mark.unit

ZONE = ZoneInfo("America/Bogota")
SEPTEMBER = DateRange(date(2026, 8, 31), date(2026, 9, 2))
DAYS = (date(2026, 8, 31), date(2026, 9, 1), date(2026, 9, 2))


# =============================================================================
# Fakes
# =============================================================================
def _area_energy(resolution: str = "day", area: str = "F&A") -> AreaEnergy:
    """1,000 kWh a day, the second day frozen; the months clipped to the range."""
    days = tuple(
        EnergySpan(day, day, None if day == DAYS[1] else Decimal(1000), Decimal(1), 1)
        for day in DAYS
    )
    months = (
        EnergySpan(DAYS[0], DAYS[0], Decimal(1000), Decimal(1), 1),
        EnergySpan(DAYS[1], DAYS[2], Decimal(2000), Decimal("0.5"), 1),
    )
    return AreaEnergy(
        area=area,
        source="export.xlsx",
        resolution="day" if resolution == "day" else "month",
        period=SEPTEMBER,
        meters=("R1",),
        days=days if resolution == "day" else (),
        months=months,
        total=EnergySpan(DAYS[0], DAYS[2], Decimal(3000), Decimal("0.67"), 2),
    )


class FakeSource:
    """An energy source with fixed answers."""

    def __init__(self, *areas: str, resolution: str = "day") -> None:
        self._areas = areas or ("F&A",)
        self._resolution = resolution
        self.asked: list[tuple[str, DateRange]] = []

    @property
    def areas(self) -> tuple[str, ...]:
        return self._areas

    def energy(self, area: str, period: DateRange) -> AreaEnergy:
        self.asked.append((area, period))
        return _area_energy(self._resolution, area)


class FakeEquivalents:
    """EQ of FORMACION and DECORADO: 50 + 50 a day."""

    def __init__(self, *, equal: bool = True) -> None:
        self.calls: list[dict[str, Any]] = []
        self._equal = equal

    def series(
        self,
        period: DateRange,
        *,
        frequency: Frequency = "day",
        procesos: Sequence[str] = (),
        areas: Sequence[str] = (),
    ) -> BatteryEquivalents:
        self.calls.append(
            {"period": period, "frequency": frequency, "procesos": procesos, "areas": areas}
        )
        rows = tuple(
            SeriesRow(day, 1, proceso, "F&A", cuenta, Decimal(1), Decimal(50), Decimal(50))
            for day in DAYS
            for proceso, cuenta in (("FORMACION", "POS_ACT_MTL"), ("DECORADO", "COMP_DEC"))
        )
        month = date(2026, 9, 1)
        return BatteryEquivalents(
            period=period,
            frequency=frequency,
            procesos=("FORMACION", "DECORADO"),
            filtered=True,
            versions=(VersionMonth(month, 1, "STD FY26"),),
            rows=rows,
            reconciliation=(
                ReconciledAmount(
                    month,
                    1,
                    "COMP_DEC",
                    Decimal(1),
                    Decimal(1),
                    Decimal(50),
                    Decimal(50 if self._equal else 49),
                    equal=self._equal,
                ),
            ),
        )


# =============================================================================
# The service
# =============================================================================
def test_the_ports_are_what_the_real_objects_satisfy() -> None:
    """BatterySeries is the EQ reader as it is; the SCADA source is an energy source."""
    assert isinstance(BatterySeries(object(), object()), EquivalentsReader)  # type: ignore[arg-type]
    assert isinstance(ScadaCounterSource(()), AreaEnergySource)


def test_each_day_is_set_against_the_eq_of_its_area() -> None:
    equivalents = FakeEquivalents()
    cross = EnergyCross([FakeSource()], equivalents)

    series = cross.cross(SEPTEMBER)

    (area,) = series.areas
    assert equivalents.calls == [
        {"period": SEPTEMBER, "frequency": "day", "procesos": (), "areas": ("F&A",)}
    ]
    assert area.procesos == ("FORMACION", "DECORADO")
    assert [(row.start, row.kwh, row.eq, row.kwh_per_eq) for row in area.rows] == [
        (DAYS[0], Decimal(1000), Decimal(100), Decimal(10)),
        (DAYS[1], None, Decimal(100), None),
        (DAYS[2], Decimal(1000), Decimal(100), Decimal(10)),
    ]
    assert (area.total.kwh, area.total.eq, area.total.kwh_per_eq) == (
        Decimal(3000),
        Decimal(300),
        Decimal(10),
    )
    assert area.reconciled_months == (date(2026, 9, 1),)
    assert series.eq_mismatches == 0


def test_a_month_cut_by_the_range_gets_the_eq_of_its_days_only() -> None:
    """The EQ is always read by day, so 2026-08 has one day of it and 2026-09 two."""
    equivalents = FakeEquivalents()

    series = EnergyCross([FakeSource()], equivalents).cross(SEPTEMBER, frequency="month")

    (area,) = series.areas
    assert equivalents.calls[0]["frequency"] == "day"
    assert [(row.start, row.end, row.kwh, row.eq) for row in area.rows] == [
        (DAYS[0], DAYS[0], Decimal(1000), Decimal(100)),
        (DAYS[1], DAYS[2], Decimal(2000), Decimal(200)),
    ]


def test_a_monthly_source_answers_only_the_monthly_cross() -> None:
    """The port takes a monthly workbook as it is; the daily cross refuses it."""
    cross = EnergyCross([FakeSource("OREC", resolution="month")], FakeEquivalents())

    assert len(cross.cross(SEPTEMBER, frequency="month").areas[0].rows) == 2
    with pytest.raises(InvalidPeriodError, match="only monthly energy"):
        cross.cross(SEPTEMBER, frequency="day")


def test_areas_are_matched_ignoring_case_and_kept_in_the_order_asked() -> None:
    cross = EnergyCross([FakeSource("F&A"), FakeSource("OREC")], FakeEquivalents())

    series = cross.cross(SEPTEMBER, areas=["orec", "f&a", "OREC"])

    assert cross.areas == ("F&A", "OREC")
    assert [area.area for area in series.areas] == ["OREC", "F&A"]


def test_every_area_with_energy_is_crossed_when_none_is_named() -> None:
    cross = EnergyCross([FakeSource("F&A", "OREC")], FakeEquivalents())

    assert [area.area for area in cross.cross(SEPTEMBER).areas] == ["F&A", "OREC"]


@pytest.mark.parametrize(
    ("options", "message"),
    [
        ({"areas": ["RC"]}, "no energy source for this area"),
        ({"frequency": "week"}, "the frequency is day or month"),
    ],
)
def test_a_question_without_an_answer_is_refused(options: dict[str, Any], message: str) -> None:
    """RC has energy but no EQ, so no source is given for it."""
    with pytest.raises(InvalidPeriodError, match=message):
        EnergyCross([FakeSource()], FakeEquivalents()).cross(SEPTEMBER, **options)


def test_two_sources_cannot_claim_one_area() -> None:
    with pytest.raises(ConfigurationError):
        EnergyCross([FakeSource("F&A"), FakeSource("F&A")], FakeEquivalents())


def test_daily_eq_that_does_not_add_up_is_reported() -> None:
    series = EnergyCross([FakeSource()], FakeEquivalents(equal=False)).cross(SEPTEMBER)

    assert series.eq_mismatches == 1


# =============================================================================
# The SCADA export
# =============================================================================
FIRST = datetime(2026, 7, 1, 19, 24, 49)  # noqa: DTZ001 - the export writes naive local time
METERS = ("RECTIFICADOR 1", "RECTIFICADOR 2", "RECTIFICADOR 3")


def _export(path: Path, hours: int = 60, *, header: Sequence[str] | None = None) -> Path:
    """An export of steady rectifiers: 100, 200 and 50 kWh an hour, counters in Wh."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dataset 1"
    sheet.append(list(header or ("t_stamp", *METERS)))
    for hour in range(hours + 1):
        sheet.append(
            [
                FIRST + timedelta(hours=hour),
                22_747_664_384.0 + hour * 100_000,
                23_587_129_344.0 + hour * 200_000,
                20_015_581_184.0 + hour * 50_000,
            ]
        )
    sheet.append([None, None, None, None])  # a blank row at the end
    buffer = io.BytesIO()
    workbook.save(buffer)
    path.write_bytes(buffer.getvalue())
    return path


def _source(path: Path) -> ScadaCounterSource:
    return ScadaCounterSource([CounterWorkbook("F&A", path, "Dataset 1", "t_stamp", METERS, ZONE)])


def test_the_export_is_read_as_local_time_and_wh(tmp_path: Path) -> None:
    source = _source(_export(tmp_path / "scada.xlsx"))

    energy = source.energy("F&A", DateRange(date(2026, 7, 2), date(2026, 7, 2)))
    counters = source.counters("F&A")

    assert source.areas == ("F&A",)
    assert counters.first_read == FIRST.replace(tzinfo=ZONE).astimezone(UTC)
    assert len(counters.read_at) == 61
    assert [(span.start, span.kwh) for span in energy.days] == [
        (date(2026, 7, 2), Decimal(24 * 350))
    ]
    assert energy.source == "scada.xlsx"
    assert energy.meters == METERS
    assert source.counters("F&A") is counters  # read once


def test_an_area_without_an_export_is_refused(tmp_path: Path) -> None:
    with pytest.raises(InvalidPeriodError):
        _source(tmp_path / "scada.xlsx").energy("OREC", SEPTEMBER)


def test_a_missing_export_is_reported(tmp_path: Path) -> None:
    with pytest.raises(SourceUnavailableError):
        _source(tmp_path / "absent.xlsx").energy("F&A", SEPTEMBER)


def test_an_export_without_a_meter_column_is_refused(tmp_path: Path) -> None:
    path = _export(tmp_path / "scada.xlsx", header=("t_stamp", "R1", "R2", "R3"))

    with pytest.raises(UnsupportedFileFormatError, match="lacks a column"):
        _source(path).energy("F&A", SEPTEMBER)


def test_a_timestamp_that_is_not_a_date_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "scada.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dataset 1"
    sheet.append(["t_stamp", *METERS])
    sheet.append(["ayer", 1.0, 2.0, 3.0])
    workbook.save(path)

    with pytest.raises(UnsupportedFileFormatError, match="not a date and time"):
        _source(path).energy("F&A", SEPTEMBER)


def test_two_exports_cannot_claim_one_area(tmp_path: Path) -> None:
    workbook = CounterWorkbook("F&A", tmp_path / "a.xlsx", "Dataset 1", "t_stamp", METERS, ZONE)

    with pytest.raises(ConfigurationError):
        ScadaCounterSource([workbook, replace(workbook, path=tmp_path / "b.xlsx")])
