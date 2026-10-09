"""Tests for the calamine workbook reader.

Layer: tests (unit)

The workbooks are built in memory with openpyxl, so the test needs no binary
fixture. What matters is that the reader reports real Excel coordinates and
hands the cells over untouched: turning them into text is not its job.
"""

from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook

from energy.adapters.calamine_workbook import CalamineWorkbookReader
from energy.domain.exceptions import EntityNotFoundError, UnsupportedFileFormatError
from energy.domain.landing import landing_rows
from energy.ports.workbook_reader import OpenWorkbook, WorkbookReader

pytestmark = pytest.mark.unit


def _workbook_bytes() -> bytes:
    """Build a small workbook with a normal sheet, an offset one and a hidden one."""
    workbook = Workbook()

    datos = workbook.active
    datos.title = "Datos"
    datos.append(["Número Artículo", "Cantidad", "Fecha Efectiva", "Activo", "Notas"])
    datos.append(["21500000A", 60, date(2025, 12, 10), True, None])
    datos.append(["111000002", 0.0636, datetime(2025, 12, 10, 11, 34, 5), False, "N/A"])  # noqa: DTZ001 - Excel has no time zones

    offset = workbook.create_sheet("Hoja1")
    offset["B3"] = "Planta "
    offset["C3"] = "Línea"
    offset["B4"] = "RC"
    offset["C4"] = "RCREFI02"

    hidden = workbook.create_sheet("BAT TERM")
    hidden.sheet_state = "hidden"
    hidden.append(["Artículo Matriz"])

    workbook.create_sheet("Vacía")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def data() -> bytes:
    """Return the bytes of the test workbook."""
    return _workbook_bytes()


def test_it_satisfies_the_ports(data: bytes) -> None:
    """The reader and what it opens are structurally the ports' protocols."""
    reader = CalamineWorkbookReader()

    assert isinstance(reader, WorkbookReader)
    assert isinstance(reader.open_workbook(data), OpenWorkbook)


def test_sheet_names_include_hidden_sheets_in_order(data: bytes) -> None:
    """A hidden sheet is still a sheet: BAT TERM is hidden in the real workbook."""
    workbook = CalamineWorkbookReader().open_workbook(data)

    assert workbook.sheet_names == ("Datos", "Hoja1", "BAT TERM", "Vacía")


def test_cells_arrive_untouched(data: bytes) -> None:
    """Numbers stay numbers and dates stay dates; nothing is turned into text here."""
    sheet = CalamineWorkbookReader().open_workbook(data).sheet("Datos")

    rows = [list(row) for row in sheet.rows]

    assert (sheet.first_row, sheet.first_column) == (1, 0)
    assert rows[0] == ["Número Artículo", "Cantidad", "Fecha Efectiva", "Activo", "Notas"]
    assert rows[1] == ["21500000A", 60.0, date(2025, 12, 10), True, ""]
    assert rows[2] == ["111000002", 0.0636, datetime(2025, 12, 10, 11, 34, 5), False, "N/A"]  # noqa: DTZ001 - Excel has no time zones


def test_a_sheet_that_starts_late_keeps_its_excel_coordinates(data: bytes) -> None:
    """The data of Hoja1 starts at B3: rows count from 1, columns from B.

    calamine pads the empty rows above the data and drops the empty column to
    its left. The reader has to report exactly that, or every row number and
    column letter of such a sheet would be off.
    """
    sheet = CalamineWorkbookReader().open_workbook(data).sheet("Hoja1")

    assert (sheet.first_row, sheet.first_column) == (1, 1)
    assert [list(row) for row in sheet.rows] == [
        ["", ""],
        ["", ""],
        ["Planta ", "Línea"],
        ["RC", "RCREFI02"],
    ]


def test_reader_and_landing_rules_agree_on_row_numbers(data: bytes) -> None:
    """Read through the landing rules, the offset sheet lands with Excel's row number."""
    sheet = CalamineWorkbookReader().open_workbook(data).sheet("Hoja1")

    landed = list(
        landing_rows(
            sheet.rows,
            first_row=sheet.first_row,
            first_column=sheet.first_column,
            header_row=3,
        )
    )

    assert [row.source_row for row in landed] == [4]
    assert landed[0].payload_json == '{"Línea":"RCREFI02","Planta":"RC"}'


def test_an_empty_sheet_has_no_rows(data: bytes) -> None:
    """A sheet with nothing in it is read as nothing, not as an error."""
    sheet = CalamineWorkbookReader().open_workbook(data).sheet("Vacía")

    assert list(sheet.rows) == []


def test_a_missing_sheet_is_reported(data: bytes) -> None:
    """Asking for a sheet the workbook lacks names the ones it has."""
    workbook = CalamineWorkbookReader().open_workbook(data)

    with pytest.raises(EntityNotFoundError) as raised:
        workbook.sheet("No existe")

    assert "Datos" in raised.value.details["available"]


def test_bytes_that_are_not_a_workbook_are_refused() -> None:
    """A renamed text file fails here, with a clear error, not deep in a load."""
    with pytest.raises(UnsupportedFileFormatError):
        CalamineWorkbookReader().open_workbook(b"this is not a spreadsheet")
