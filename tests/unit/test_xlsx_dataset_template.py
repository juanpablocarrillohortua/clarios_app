"""The template of a dataset: exactly the spec's headers, read back as Bronze reads them.

Every template is opened with the reader the Bronze loader uses, and its header
row keyed by the Bronze rules. The keys must be the spec's, and the sheet must
hold nothing else: the formats and drop-down lists that help people type live
on the columns, never in cells.
"""

from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING, Any

import pytest
from openpyxl import load_workbook

from energy.adapters.calamine_workbook import CalamineWorkbookReader
from energy.adapters.xlsx_dataset_template import XLSX_MEDIA_TYPE, XlsxDatasetTemplateWriter
from energy.adapters.yaml_dataset_specs import load_dataset_specs
from energy.domain.dataset_spec import DatasetSpec, parse_dataset_spec
from energy.domain.landing import cell_text, header_keys

if TYPE_CHECKING:
    from pathlib import Path


def made_up(header_row: int = 1, allowed: list[str] | None = None) -> DatasetSpec:
    """A spec with a repeated header, a date, a whole number and a list."""
    return parse_dataset_spec(
        {
            "version": 1,
            "key": "example",
            "kind": "snapshot",
            "label": "datasets.example.name",
            "description": "datasets.example.description",
            "entry_modes": ["upload"],
            "bronze": {
                "source_key": "manual_sheets",
                "sheet": "Hoja de prueba",
                "dataset_key": "example",
                "header_row": header_row,
            },
            "silver": {"table": "silver.example"},
            "natural_key": ["Código"],
            "columns": [
                {"header": "Código", "type": "text", "required": True},
                {"header": "Fecha", "type": "date"},
                {"header": "Cantidad", "type": "integer"},
                {"header": "Código", "type": "text"},
                {"header": "Activo", "type": "text", "allowed": allowed or ["Sí", "No"]},
            ],
        },
        "made up",
    )


def read_back(data: bytes, spec: DatasetSpec) -> list[list[Any]]:
    """The template's rows, as the Bronze loader reads them; asserts its one sheet."""
    workbook = CalamineWorkbookReader().open_workbook(data)
    assert workbook.sheet_names == (spec.bronze.sheet,)
    sheet = workbook.sheet(spec.bronze.sheet)
    assert sheet.first_column == 0
    return [list(row) for row in sheet.rows]


def specs(repo_root: Path) -> list[DatasetSpec]:
    return [*load_dataset_specs(repo_root / "config"), made_up(), made_up(header_row=3)]


def test_every_template_holds_exactly_the_specs_headers(repo_root: Path) -> None:
    writer = XlsxDatasetTemplateWriter()
    for spec in specs(repo_root):
        rows = read_back(writer.write(spec), spec)
        header_index = spec.bronze.header_row - 1

        assert len(rows) == spec.bronze.header_row, f"{spec.key}: the sheet holds data"
        assert header_keys(rows[header_index]) == spec.keys, spec.key
        assert tuple(cell_text(cell) for cell in rows[header_index]) == spec.headers
        assert all(cell_text(cell) is None for row in rows[:header_index] for cell in row)


def test_a_repeated_header_is_written_twice_and_read_as_bronze_keys_it() -> None:
    spec = made_up()

    rows = read_back(XlsxDatasetTemplateWriter().write(spec), spec)

    assert rows[0] == ["Código", "Fecha", "Cantidad", "Código", "Activo"]
    assert header_keys(rows[0]) == ("Código", "Fecha", "Cantidad", "Código #2", "Activo")


def test_each_column_is_formatted_for_its_type_and_the_header_stays_in_view() -> None:
    data = XlsxDatasetTemplateWriter().write(made_up(header_row=3))
    sheet = load_workbook(BytesIO(data)).active

    formats = [sheet.column_dimensions[letter].number_format for letter in "ABCDE"]
    assert formats == ["@", "yyyy-mm-dd", "0", "@", "@"]
    assert sheet.freeze_panes == "A4"
    assert all(sheet.cell(row=3, column=column).font.bold for column in range(1, 6))


def test_allowed_values_are_offered_in_a_list_below_the_header() -> None:
    data = XlsxDatasetTemplateWriter().write(made_up(header_row=3))
    sheet = load_workbook(BytesIO(data)).active

    (validation,) = sheet.data_validations.dataValidation
    assert validation.type == "list"
    assert validation.formula1 == '"Sí,No"'
    assert str(validation.sqref) == "E4:E1048576"


@pytest.mark.parametrize("allowed", [["Sí, claro", "No"], ['El "otro"', "No"], ["x" * 300]])
def test_values_a_list_cannot_hold_are_not_offered(allowed: list[str]) -> None:
    data = XlsxDatasetTemplateWriter().write(made_up(allowed=allowed))
    sheet = load_workbook(BytesIO(data)).active

    assert sheet.data_validations.dataValidation == []


def test_the_writer_names_its_file_type() -> None:
    writer = XlsxDatasetTemplateWriter()

    assert writer.media_type == XLSX_MEDIA_TYPE
    assert writer.suffix == ".xlsx"
