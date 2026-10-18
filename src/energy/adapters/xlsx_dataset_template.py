"""Write the template of a dataset as an ``.xlsx`` workbook.

Layer: adapters

Responsibility:
    Implements :class:`~energy.ports.dataset_template.DatasetTemplateWriter`
    with ``openpyxl``. The workbook has one sheet, named as the spec names it,
    with the spec's headers on its header row, in order, repeats included, and
    no data.

    It also helps the person who fills it in, without adding a cell:

    - each column is formatted for its type: text columns as Text, so Excel
      does not turn an item code such as ``231001483`` into a number and then
      into ``2.31E+08``; date columns as dates; whole numbers without decimals;
    - a column with allowed values offers them in a drop-down list;
    - the header row is bold and stays in view when scrolling.

    Formats and lists live on the columns, not in cells, so the sheet still
    holds the header row and nothing else: read back by the Bronze rules, it
    gives exactly the spec's payload keys. A test reads every template that way.

    There is no sheet of instructions: its text would be Spanish written in
    Python, which the language rule forbids, and the web app explains instead.

Depends on:
    ``energy.domain``, ``openpyxl``, standard library.
"""

from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING, Any, Final

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.worksheet.datavalidation import DataValidation  # type: ignore[import-untyped]

from energy.domain.dataset_spec import ColumnType

if TYPE_CHECKING:
    from energy.domain.dataset_spec import ColumnSpec, DatasetSpec

XLSX_MEDIA_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Excel's built-in formats: "@" is Text, "0" a whole number.
_NUMBER_FORMATS: Final = {
    ColumnType.TEXT: "@",
    ColumnType.INTEGER: "0",
    ColumnType.DECIMAL: "General",
    ColumnType.DATE: "yyyy-mm-dd",
}
_LAST_ROW: Final = 1_048_576  # Excel's last row
# Excel writes a list's values into one formula of at most 255 characters,
# separated by commas and quoted, so a value may hold neither.
_LIST_FORMULA_MAX: Final = 255
_MIN_WIDTH: Final = 10
_WIDTH_MARGIN: Final = 2


def _list_formula(values: tuple[str, ...]) -> str | None:
    """Return the formula of a drop-down list, or ``None`` if Excel cannot hold it."""
    if any("," in value or '"' in value for value in values):
        return None
    formula = '"' + ",".join(values) + '"'
    return formula if len(formula) <= _LIST_FORMULA_MAX else None


class XlsxDatasetTemplateWriter:
    """Writes dataset templates as ``.xlsx`` workbooks."""

    media_type: str = XLSX_MEDIA_TYPE
    suffix: str = ".xlsx"

    def write(self, spec: DatasetSpec) -> bytes:
        """Return the template of one dataset.

        Args:
            spec: the dataset's spec.

        Returns:
            The workbook's bytes.
        """
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = spec.bronze.sheet
        header_row = spec.bronze.header_row
        bold = Font(bold=True)
        for index, column in enumerate(spec.columns, start=1):
            cell = sheet.cell(row=header_row, column=index, value=column.header)
            cell.font = bold
            self._format_column(sheet, index, column, header_row)
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _format_column(sheet: Any, index: int, column: ColumnSpec, header_row: int) -> None:
        """Format one column for its type, and offer its allowed values.

        Args:
            sheet: the openpyxl worksheet, which has no type information.
            index: the column's 1-based position.
            column: the column's spec.
            header_row: the row of the headers; the formats start below it.
        """
        letter = get_column_letter(index)
        dimension = sheet.column_dimensions[letter]
        dimension.number_format = _NUMBER_FORMATS[column.type]
        dimension.width = max(_MIN_WIDTH, len(column.header) + _WIDTH_MARGIN)
        formula = _list_formula(column.allowed) if column.allowed else None
        if formula is not None:
            validation = DataValidation(type="list", formula1=formula, allow_blank=True)
            validation.add(f"{letter}{header_row + 1}:{letter}{_LAST_ROW}")
            sheet.add_data_validation(validation)


__all__ = ["XLSX_MEDIA_TYPE", "XlsxDatasetTemplateWriter"]
